from flask import Flask, render_template, jsonify, send_file, abort, request, session, redirect
from functools import wraps
import os
import re
import hmac
import secrets
import subprocess
import zipfile

from pathlib import Path
from datetime import datetime, timedelta
import json
import wave
import io
import time
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

BASE_DIR = Path("/opt/toyota-malaysia-ai")
REPORT_DIR = BASE_DIR / "reports" / "calls"
RECORDING_DIR = BASE_DIR / "recordings"
ACTIVE_FILE = Path("/run/toyota-ai/active-calls.json")

app = Flask(__name__)

SESSION_SECRET_FILE = BASE_DIR / "secrets" / "dashboard.secret"
SESSION_SECRET_FILE.parent.mkdir(parents=True, exist_ok=True)

if SESSION_SECRET_FILE.exists():
    _session_secret = SESSION_SECRET_FILE.read_text().strip()
else:
    _session_secret = secrets.token_hex(32)
    SESSION_SECRET_FILE.write_text(_session_secret)
    try:
        SESSION_SECRET_FILE.chmod(0o600)
    except Exception:
        pass

app.secret_key = _session_secret
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=8)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"

TOYOTA_REPORT_PREFIX = "/toyota-reports"




# ===== TOYOTA VICIDIAL AUTH =====

def _read_vicidial_config():
    """
    Read VICIdial DB settings from /etc/astguiclient.conf.
    Supports '=' and '=>'.
    Values are never returned to the browser.
    """
    config_file = Path("/etc/astguiclient.conf")
    values = {}

    if not config_file.exists():
        raise RuntimeError("VICIdial configuration file not found")

    text = config_file.read_text(errors="ignore")

    for key in (
        "VARDB_server",
        "VARDB_database",
        "VARDB_user",
        "VARDB_pass",
        "VARDB_port",
    ):
        pattern = rf"^\s*{re.escape(key)}\s*(?:=>|=)\s*[\"']?([^\"'\s;]+)"
        match = re.search(pattern, text, re.MULTILINE)
        if match:
            values[key] = match.group(1)

    # Fallback defaults used by VICIdial.
    values.setdefault("VARDB_server", "localhost")
    values.setdefault("VARDB_database", "asterisk")
    values.setdefault("VARDB_user", "cron")
    values.setdefault("VARDB_port", "3306")

    if not values.get("VARDB_pass"):
        raise RuntimeError("VICIdial DB password not found")

    return values


def _vicidial_auth(username, password):
    """
    Authenticate against the same VICIdial credentials.

    Dashboard policy:
      active = Y
      user_level = 9 ONLY

    Password verification follows VICIdial's pass_hash_enabled
    setting and bp.pl hashing mechanism.
    """
    username = str(username or "").strip()
    password = str(password or "")

    if not username or not password:
        return False

    # Prevent SQL/user-name abuse before querying.
    if not re.fullmatch(r"[A-Za-z0-9_.@-]{1,50}", username):
        return False

    try:
        import pymysql
        from pymysql.cursors import DictCursor

        cfg = _read_vicidial_config()

        conn = pymysql.connect(
            host=cfg["VARDB_server"],
            port=int(cfg.get("VARDB_port") or 3306),
            user=cfg["VARDB_user"],
            password=cfg["VARDB_pass"],
            database=cfg["VARDB_database"],
            cursorclass=DictCursor,
            connect_timeout=5,
            read_timeout=5,
            write_timeout=5,
            charset="utf8mb4",
        )

        try:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT
                        user,
                        pass,
                        pass_hash,
                        user_level,
                        active,
                        failed_login_count,
                        last_login_date
                    FROM vicidial_users
                    WHERE user=%s
                    LIMIT 1
                    """,
                    (username,),
                )
                user = cur.fetchone()

                if not user:
                    return False

                # STRICT Toyota dashboard policy.
                if str(user.get("active") or "").upper() != "Y":
                    return False

                if int(user.get("user_level") or 0) != 9:
                    return False

                cur.execute(
                    """
                    SELECT pass_hash_enabled
                    FROM system_settings
                    LIMIT 1
                    """
                )
                settings = cur.fetchone() or {}

                pass_hash_enabled = int(
                    settings.get("pass_hash_enabled") or 0
                )

                if pass_hash_enabled > 0:
                    bp = Path("/var/www/html/agc/bp.pl")

                    if not bp.exists():
                        return False

                    # This follows VICIdial's own authentication mechanism.
                    proc = subprocess.run(
                        [
                            str(bp),
                            "--pass=" + password,
                        ],
                        stdout=subprocess.PIPE,
                        stderr=subprocess.DEVNULL,
                        text=True,
                        timeout=5,
                    )

                    generated_hash = re.sub(
                        r"PHASH:|\s+",
                        "",
                        proc.stdout or "",
                    )

                    stored_hash = str(
                        user.get("pass_hash") or ""
                    ).strip()

                    if not generated_hash or not stored_hash:
                        return False

                    valid = hmac.compare_digest(
                        generated_hash,
                        stored_hash,
                    )
                else:
                    stored_password = str(
                        user.get("pass") or ""
                    )

                    valid = hmac.compare_digest(
                        password,
                        stored_password,
                    )

                if not valid:
                    return False

                # Successful login: reset VICIdial failed counter
                # and update login time, same general behavior as VICIdial.
                try:
                    cur.execute(
                        """
                        UPDATE vicidial_users
                        SET
                            last_login_date=NOW(),
                            last_ip=%s,
                            failed_login_count=0
                        WHERE user=%s
                        """,
                        (
                            request.remote_addr or "",
                            username,
                        ),
                    )
                    conn.commit()
                except Exception:
                    conn.rollback()

                return True

        finally:
            conn.close()

    except Exception:
        return False


def _login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("vicidial_admin_authenticated"):
            if request.path.startswith("/api/"):
                return jsonify({
                    "error": "authentication_required"
                }), 401

            return redirect(
                TOYOTA_REPORT_PREFIX + "/login"
            )

        return view(*args, **kwargs)

    return wrapped


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        if session.get("vicidial_admin_authenticated"):
            return redirect(TOYOTA_REPORT_PREFIX + "/")

        return render_template("login.html")

    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")

    if _vicidial_auth(username, password):
        session.clear()
        session.permanent = True
        session["vicidial_admin_authenticated"] = True
        session["vicidial_username"] = username

        return redirect(TOYOTA_REPORT_PREFIX + "/")

    return render_template(
        "login.html",
        error="Invalid VICIdial username or password."
    ), 401


@app.route("/logout")
def logout():
    session.clear()
    return redirect(TOYOTA_REPORT_PREFIX + "/login")



def safe_recording_path(filename):
    requested = (RECORDING_DIR / filename).resolve()
    try:
        requested.relative_to(RECORDING_DIR.resolve())
    except ValueError:
        abort(403)
    if not requested.is_file():
        abort(404)
    return requested


def get_recording_duration(recording_path):
    if not recording_path:
        return None
    try:
        path = Path(recording_path).resolve()
        path.relative_to(RECORDING_DIR.resolve())
        if not path.is_file():
            return None
        with wave.open(str(path), "rb") as wf:
            frames = wf.getnframes()
            rate = wf.getframerate()
            return round(frames / rate, 2) if rate > 0 else None
    except Exception:
        return None


def load_json_report(path):
    try:
        with path.open("r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, dict):
            return None
        data["_report_file"] = path.name
        data["recording_duration_seconds"] = get_recording_duration(data.get("recording"))
        return data
    except Exception:
        return None


def load_reports():
    reports = []
    if not REPORT_DIR.exists():
        return reports

    for path in REPORT_DIR.glob("*.json"):
        data = load_json_report(path)
        if data:
            reports.append(data)

    grouped = {}
    for report in reports:
        call_id = str(report.get("call_id") or report.get("_report_file"))
        existing = grouped.get(call_id)
        if existing is None:
            grouped[call_id] = report
            continue

        existing_finished = bool(existing.get("end_time"))
        current_finished = bool(report.get("end_time"))
        if current_finished and not existing_finished:
            grouped[call_id] = report
        elif current_finished == existing_finished:
            existing_time = existing.get("end_time") or existing.get("start_time") or ""
            current_time = current_time = report.get("end_time") or report.get("start_time") or ""
            if current_time >= existing_time:
                grouped[call_id] = report

    reports = list(grouped.values())
    reports.sort(key=lambda x: x.get("start_time") or "", reverse=True)
    return reports


def calculate_stats(reports):
    total = len(reports)

    phone_leads = sum(
        1 for r in reports
        if r.get("disposition") in ("Interested – Agent Call", "Interested – Specific Model")
    )
    whatsapp_leads = sum(
        1 for r in reports
        if r.get("disposition") == "Interested – WhatsApp"
    )
    interested = sum(1 for r in reports if r.get("intent") == "Interested")
    not_interested = sum(1 for r in reports if r.get("disposition") == "Not Interested")
    callback = sum(1 for r in reports if r.get("disposition") == "Call Back Later")
    wrong_number = sum(1 for r in reports if r.get("disposition") == "Wrong Number")
    dnc = sum(1 for r in reports if r.get("disposition") == "Do Not Contact")

    hangup = sum(
        1 for r in reports
        if "customer hang up" in str(r.get("conversation_status") or r.get("status") or "").lower()
        or "customer hang up" in str(r.get("disposition") or "").lower()
    )

    total_duration = sum(float(r.get("duration_seconds") or 0) for r in reports)
    total_recording_duration = sum(float(r.get("recording_duration_seconds") or 0) for r in reports)
    recording_count = sum(1 for r in reports if r.get("recording_duration_seconds") is not None)

    return {
        "total": total,
        "phone_leads": phone_leads,
        "whatsapp_leads": whatsapp_leads,
        "interested": interested,
        "not_interested": not_interested,
        "callback": callback,
        "wrong_number": wrong_number,
        "dnc": dnc,
        "hangup": hangup,
        "total_duration": round(total_duration, 2),
        "total_recording_duration": round(total_recording_duration, 2),
        "recording_count": recording_count,
    }


def get_live_calls():
    if not ACTIVE_FILE.exists():
        return []
    try:
        data = json.loads(ACTIVE_FILE.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            return []
        live = []
        for call_id, call in data.items():
            if not isinstance(call, dict):
                continue
            if str(call.get("stage", "")).lower() == "completed":
                continue
            item = dict(call)
            item.setdefault("call_id", call_id)
            live.append(item)
        return live
    except Exception:
        return []



EXCEL_HEADERS = [
    "Call_ID",
    "Call_Date",
    "Call_Start_Time",
    "Call_End_Time",
    "Call_Duration_Sec",
    "Lead_ID",
    "Phone_Number",
    "Customer_Name",
    "Conversation_Status",
    "Disposition",
    "Customer_Captured_Model",
    "Car_Type",
    "Purchase_Type",
    "Current_Car",
    "Interested_Model",
    "Budget",
    "Monthly_Payment",
    "Purchase_Timeline",
    "Callback_Time",
    "Contact_Preference",
    "Customer_Language",
]


def _conversation_parts(report):
    conversation = report.get("conversation") or []

    customer = []
    ai = []
    languages = []

    for item in conversation:
        if not isinstance(item, dict):
            continue

        speaker = str(item.get("speaker") or "").lower()

        value = (
            item.get("english_text")
            or item.get("original_text")
            or item.get("text")
            or ""
        ).strip()

        if speaker == "customer":
            if value:
                customer.append(value)

            if item.get("language"):
                languages.append(str(item["language"]))

        elif speaker == "ai":
            if value:
                ai.append(value)

    kb_questions = []
    kb_answers = []
    kb_models = []
    kb_variants = []

    for item in report.get("kb_history") or []:
        if not isinstance(item, dict):
            continue

        question = str(item.get("question") or "").strip()
        answer = str(item.get("answer") or "").strip()

        if question:
            kb_questions.append(question)

        if answer:
            kb_answers.append(answer)

        if item.get("model"):
            kb_models.append(str(item["model"]))

        if item.get("variant"):
            kb_variants.append(str(item["variant"]))

    return (
        "\n".join(customer),
        "\n".join(ai),
        ", ".join(dict.fromkeys(languages)),
        "\n".join(kb_questions),
        "\n".join(kb_answers),
        ", ".join(dict.fromkeys(kb_models)),
        ", ".join(dict.fromkeys(kb_variants)),
    )


def _excel_safe(value):
    """Remove characters that Excel/openpyxl cannot store in worksheet cells."""
    if value is None:
        return None

    if isinstance(value, str):
        return "".join(
            ch for ch in value
            if ch in "\\t\\n\\r"
            or 0x20 <= ord(ch) <= 0xD7FF
            or 0xE000 <= ord(ch) <= 0xFFFD
            or 0x10000 <= ord(ch) <= 0x10FFFF
        )

    if isinstance(value, (list, tuple)):
        return ", ".join(str(_excel_safe(v)) for v in value)

    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)

    return value


def _report_excel_row(report):
    start = report.get("start_time") or ""
    end = report.get("end_time") or ""

    call_date = (
        start[:10]
        if start
        else (end[:10] if end else "")
    )

    disposition = report.get("disposition") or ""

    status = (
        report.get("conversation_status")
        or report.get("status")
        or (
            "Customer Hang Up"
            if disposition.lower() == "customer hang up"
            else ("Completed" if end else "Conversation Not Done")
        )
    )

    return [
        report.get("call_id"),
        call_date,
        start,
        end,
        report.get("duration_seconds"),
        report.get("lead_id"),
        report.get("phone_number"),
        report.get("first_name"),
        status,
        disposition,

        report.get("customer_captured_model"),
        report.get("car_type"),
        report.get("purchase_type"),
        report.get("current_car"),

        (
            report.get("interested_model")
            or report.get("preferred_model")
        ),

        report.get("budget"),
        report.get("monthly_payment"),
        report.get("purchase_timeline"),
        report.get("callback_time"),

        (
            report.get("contact_preference")
            or report.get("final_contact")
        ),

        report.get("customer_language"),
    ]


@app.route("/api/export/excel")
@_login_required
def api_export_excel():
    reports = load_reports()

    date_from = request.args.get(
        "date_from",
        ""
    ).strip()

    date_to = request.args.get(
        "date_to",
        ""
    ).strip()

    phone = request.args.get(
        "phone",
        ""
    ).strip()

    df = (
        datetime.fromisoformat(date_from).date()
        if date_from
        else None
    )

    dt = (
        datetime.fromisoformat(date_to).date()
        if date_to
        else None
    )

    filtered = [
        r for r in reports
        if report_matches_filters(
            r,
            df,
            dt,
            phone,
        )
    ]

    wb = Workbook()

    ws = wb.active
    ws.title = "Call_Report"

    ws.append(EXCEL_HEADERS)

    for report in filtered:
        ws.append(
            [_excel_safe(v) for v in _report_excel_row(report)]
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = [
        38, 12, 24, 24, 16, 12, 18, 20,
        22, 24, 24, 16, 18, 24, 24, 20,
        20, 22, 22, 22, 18,
    ]

    # ---------------------------------------------------------
    # Call Summary sheet
    # ---------------------------------------------------------
    summary = wb.create_sheet("Call_Summary")

    summary.append(["Toyota Malaysia AI - Call Summary"])
    summary.append([])
    summary.append(["Metric", "Value"])

    total_calls = len(filtered)

    phone_leads = sum(
        1 for r in filtered
        if r.get("disposition") in (
            "Interested – Agent Call",
            "Interested – Specific Model",
        )
    )

    whatsapp_leads = sum(
        1 for r in filtered
        if r.get("disposition") == "Interested – WhatsApp"
    )

    interested = sum(
        1 for r in filtered
        if str(r.get("disposition") or "").startswith("Interested")
    )

    not_interested = sum(
        1 for r in filtered
        if r.get("disposition") == "Not Interested"
    )

    callback_later = sum(
        1 for r in filtered
        if r.get("disposition") == "Call Back Later"
    )

    wrong_number = sum(
        1 for r in filtered
        if r.get("disposition") == "Wrong Number"
    )

    do_not_contact = sum(
        1 for r in filtered
        if r.get("disposition") == "Do Not Contact"
    )

    customer_hangup = sum(
        1 for r in filtered
        if r.get("disposition") == "Customer Hang Up"
    )

    durations = []

    for r in filtered:
        try:
            value = float(r.get("duration_seconds") or 0)
            if value > 0:
                durations.append(value)
        except (TypeError, ValueError):
            pass

    average_duration = (
        round(sum(durations) / len(durations), 2)
        if durations else 0
    )

    summary_rows = [
        ["Total Calls", total_calls],
        ["Phone Leads", phone_leads],
        ["WhatsApp Leads", whatsapp_leads],
        ["Interested", interested],
        ["Not Interested", not_interested],
        ["Call Back Later", callback_later],
        ["Wrong Number", wrong_number],
        ["Do Not Contact", do_not_contact],
        ["Customer Hang Up", customer_hangup],
        ["Average Duration (Sec)", average_duration],
    ]

    for row in summary_rows:
        summary.append(row)

    summary.append([])
    summary.append(["Applied Filters", ""])
    summary.append(["Date From", date_from or "All"])
    summary.append(["Date To", date_to or "All"])
    summary.append(["Phone", phone or "All"])

    summary.freeze_panes = "A4"
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 24

    for cell in summary[1]:
        cell.font = cell.font.copy(bold=True, size=14)

    for cell in summary[3]:
        cell.font = cell.font.copy(bold=True)

    for index, width in enumerate(
        widths,
        start=1,
    ):
        ws.column_dimensions[
            get_column_letter(index)
        ].width = width

    for cell in ws[1]:
        cell.font = cell.font.copy(
            bold=True
        )

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return send_file(
        output,
        mimetype=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name="toyota_call_report.xlsx",
    )


def parse_date(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except Exception:
        return None


def report_matches_filters(report, date_from=None, date_to=None, phone=None):
    dt = parse_date(report.get("start_time") or report.get("end_time"))
    if date_from and (not dt or dt.date() < date_from):
        return False
    if date_to and (not dt or dt.date() > date_to):
        return False
    if phone:
        haystack = str(report.get("phone_number") or "").lower()
        if phone.lower() not in haystack:
            return False
    return True


@app.route("/")
@_login_required
def dashboard():
    return render_template("dashboard.html")


@app.route("/api/reports")
@_login_required
def api_reports():
    reports = load_reports()

    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    phone = request.args.get("phone", "").strip()

    df = datetime.fromisoformat(date_from).date() if date_from else None
    dt = datetime.fromisoformat(date_to).date() if date_to else None

    filtered = [r for r in reports if report_matches_filters(r, df, dt, phone)]

    return jsonify({
        "stats": calculate_stats(filtered),
        "reports": filtered,
        "total_available": len(reports),
        "server_time": datetime.now().astimezone().isoformat(),
    })


@app.route("/api/live-calls")
@_login_required
def api_live_calls():
    live_calls = get_live_calls()
    result = []

    for r in live_calls:
        item = dict(r)
        start_time = r.get("start_time")
        elapsed = r.get("duration_seconds") or 0

        if start_time:
            try:
                started = datetime.fromisoformat(start_time)
                elapsed = max(0, (datetime.now(started.tzinfo) - started).total_seconds())
            except Exception:
                pass

        item["live_duration_seconds"] = round(float(elapsed), 1)

        conversation = r.get("conversation") or []
        if conversation:
            last = conversation[-1]
            item["last_speaker"] = last.get("speaker")
            item["last_message"] = last.get("english_text") or last.get("original_text") or ""
        else:
            item["last_speaker"] = None
            item["last_message"] = ""

        result.append(item)

    return jsonify({
        "count": len(result),
        "calls": result,
        "server_time": datetime.now().astimezone().isoformat(),
    })


@app.route("/recordings/<path:filename>")
@_login_required
def recording(filename):
    requested = safe_recording_path(filename)
    return send_file(requested, mimetype="audio/wav", conditional=True)


@app.route("/recordings/channel/<channel>/<path:filename>")
@_login_required
def recording_channel(channel, filename):
    if channel not in ("customer", "ai", "both"):
        abort(400)

    requested = safe_recording_path(filename)

    try:
        with wave.open(str(requested), "rb") as src:
            channels = src.getnchannels()
            sample_width = src.getsampwidth()
            frame_rate = src.getframerate()
            frames = src.readframes(src.getnframes())

        if channels == 1:
            output = io.BytesIO()
            with wave.open(output, "wb") as dst:
                dst.setnchannels(1)
                dst.setsampwidth(sample_width)
                dst.setframerate(frame_rate)
                dst.writeframes(frames)
            output.seek(0)
            return send_file(output, mimetype="audio/wav",
                             download_name=f"{requested.stem}-{channel}.wav")

        if channels != 2:
            abort(415)

        bytes_per_sample = sample_width
        frame_size = channels * bytes_per_sample
        total_frames = len(frames) // frame_size
        customer = bytearray()
        ai = bytearray()

        for i in range(total_frames):
            offset = i * frame_size
            customer.extend(frames[offset:offset + bytes_per_sample])
            ai.extend(frames[offset + bytes_per_sample:offset + frame_size])

        selected = frames if channel == "both" else (bytes(customer) if channel == "customer" else bytes(ai))

        output = io.BytesIO()
        with wave.open(output, "wb") as dst:
            dst.setnchannels(2 if channel == "both" else 1)
            dst.setsampwidth(sample_width)
            dst.setframerate(frame_rate)
            dst.writeframes(selected)
        output.seek(0)

        return send_file(output, mimetype="audio/wav",
                         download_name=f"{requested.stem}-{channel}.wav")
    except Exception:
        abort(500)




# ===== TOYOTA CALL REPORT DOWNLOAD =====

@app.route("/api/download/call-reports")
@_login_required
def api_download_call_reports():
    reports = load_reports()

    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    phone = request.args.get("phone", "").strip()

    try:
        df = datetime.fromisoformat(date_from).date() if date_from else None
        dt = datetime.fromisoformat(date_to).date() if date_to else None
    except ValueError:
        return jsonify({"error": "Invalid date filter"}), 400

    filtered = [
        r for r in reports
        if report_matches_filters(r, df, dt, phone)
    ]

    output = io.BytesIO()

    with zipfile.ZipFile(
        output,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:

        for index, report in enumerate(filtered, start=1):
            filename = str(
                report.get("_report_file")
                or report.get("call_id")
                or f"call_{index}"
            )

            filename = re.sub(
                r"[^A-Za-z0-9_.-]+",
                "_",
                filename,
            )

            if not filename.endswith(".json"):
                filename += ".json"

            clean_report = dict(report)
            clean_report.pop("_report_file", None)
            clean_report.pop("recording_duration_seconds", None)

            archive.writestr(
                filename,
                json.dumps(
                    clean_report,
                    ensure_ascii=False,
                    indent=2,
                ),
            )

    output.seek(0)

    return send_file(
        output,
        mimetype="application/zip",
        as_attachment=True,
        download_name="toyota_call_reports.zip",
    )

@app.route("/health")
def health():
    reports = load_reports()
    live = get_live_calls()
    return jsonify({
        "status": "ok",
        "reports": len(reports),
        "active_calls": len(live),
        "time": datetime.now().astimezone().isoformat()
    })


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=9080, debug=False)

